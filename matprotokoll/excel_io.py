"""Skriv mätprotokoll-blad enligt mallens layout och färger."""

from __future__ import annotations

from copy import copy
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from matprotokoll.constants import (
    COVER_SHEET_NAME,
    HEADER_FILL_EXHAUST,
    HEADER_FILL_SUPPLY,
    PROTO_SHEET_CANDIDATES,
)
from matprotokoll.models import ProtocolRow
from matprotokoll.paths import template_xlsx
from matprotokoll.sheet_names import uniquify_sheet_name

TEMPLATE_PATH = template_xlsx()

_THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
_FILL_EX = PatternFill("solid", fgColor=HEADER_FILL_EXHAUST)
_FILL_SU = PatternFill("solid", fgColor=HEADER_FILL_SUPPLY)
_FONT = Font(name="Calibri", size=11)
_HEADER_FONT = Font(name="Calibri", size=10, bold=True)
_TOTAL_FONT = Font(name="Calibri", size=11, bold=True)


def default_template_path() -> Path:
    return TEMPLATE_PATH


def load_workbook_from_bytes_or_path(src: Path | str | bytes | BytesIO | None = None) -> Workbook:
    if src is None:
        return load_workbook(TEMPLATE_PATH)
    if isinstance(src, (bytes, bytearray)):
        return load_workbook(BytesIO(src))
    if isinstance(src, BytesIO):
        src.seek(0)
        return load_workbook(src)
    return load_workbook(src)


def _prototype_sheet(wb: Workbook) -> Worksheet | None:
    for name in PROTO_SHEET_CANDIDATES:
        if name in wb.sheetnames:
            return wb[name]
    for name in wb.sheetnames:
        if name != COVER_SHEET_NAME:
            return wb[name]
    return None


def ensure_cover(wb: Workbook, template: Workbook | None = None) -> None:
    if COVER_SHEET_NAME in wb.sheetnames:
        return
    src_wb = template or load_workbook(TEMPLATE_PATH)
    if COVER_SHEET_NAME not in src_wb.sheetnames:
        return
    src = src_wb[COVER_SHEET_NAME]
    dst = wb.create_sheet(COVER_SHEET_NAME, 0)
    for row in src.iter_rows():
        for cell in row:
            nc = dst.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                nc.font = copy(cell.font)
                nc.fill = copy(cell.fill)
                nc.border = copy(cell.border)
                nc.alignment = copy(cell.alignment)
                nc.number_format = cell.number_format
    for merged in src.merged_cells.ranges:
        dst.merge_cells(str(merged))
    for letter, dim in src.column_dimensions.items():
        dst.column_dimensions[letter].width = dim.width
    for idx, dim in src.row_dimensions.items():
        dst.row_dimensions[idx].height = dim.height


def _copy_header_and_layout(proto: Worksheet, dest: Worksheet) -> None:
    dest.row_dimensions[1].height = proto.row_dimensions[1].height or 39
    for col in range(1, 14):
        letter = get_column_letter(col)
        dest.column_dimensions[letter].width = proto.column_dimensions[letter].width
        src = proto.cell(1, col)
        dst = dest.cell(1, col, src.value)
        if src.has_style:
            dst.font = copy(src.font)
            dst.fill = copy(src.fill)
            dst.border = copy(src.border)
            dst.alignment = copy(src.alignment)
            dst.number_format = src.number_format
        dst.alignment = Alignment(wrap_text=True, vertical="center")


def _style_data_cell(cell, *, exhaust: bool | None) -> None:
    cell.font = _FONT
    cell.border = _THIN
    cell.alignment = Alignment(vertical="center")
    if exhaust is True:
        cell.fill = _FILL_EX
    elif exhaust is False:
        cell.fill = _FILL_SU


def write_protocol_rows(ws: Worksheet, rows: list[ProtocolRow], start_row: int = 2) -> int:
    """Skriv rader från start_row. Returnerar radnumret för totalsumman."""
    r = start_row
    for row in rows:
        ws.cell(r, 1, row.room or None)
        ws.cell(r, 1).font = _FONT
        if row.exhaust_name:
            ws.cell(r, 2, row.exhaust_name)
            ws.cell(r, 3, row.exhaust_size)
            ws.cell(r, 4, row.exhaust_behov)
            for c in range(2, 8):
                _style_data_cell(ws.cell(r, c), exhaust=True)
        if row.supply_name:
            ws.cell(r, 8, row.supply_name)
            ws.cell(r, 9, row.supply_size)
            ws.cell(r, 10, row.supply_behov)
            for c in range(8, 14):
                _style_data_cell(ws.cell(r, c), exhaust=False)
        r += 1
    last_data = r - 1
    total_row = last_data + 2 if last_data >= start_row else start_row + 1
    ws.cell(total_row, 1, "Kokonaisilmamäärät/Totalt").font = _TOTAL_FONT
    if last_data >= start_row:
        ws.cell(total_row, 4, f"=SUM(D{start_row}:D{last_data})")
        ws.cell(total_row, 5, f"=SUM(E{start_row}:E{last_data})")
        ws.cell(total_row, 10, f"=SUM(J{start_row}:J{last_data})")
        ws.cell(total_row, 11, f"=SUM(K{start_row}:K{last_data})")
    for c in (4, 5):
        _style_data_cell(ws.cell(total_row, c), exhaust=True)
        ws.cell(total_row, c).font = _TOTAL_FONT
    for c in (10, 11):
        _style_data_cell(ws.cell(total_row, c), exhaust=False)
        ws.cell(total_row, c).font = _TOTAL_FONT
    return total_row


def add_protocol_sheet(
    wb: Workbook,
    sheet_name: str,
    rows: list[ProtocolRow],
    *,
    template: Workbook | None = None,
) -> Worksheet:
    name = uniquify_sheet_name(sheet_name, wb.sheetnames)
    proto_wb = template
    if proto_wb is None:
        try:
            proto_wb = load_workbook(TEMPLATE_PATH)
        except Exception:
            proto_wb = wb
    proto = _prototype_sheet(proto_wb)
    ws = wb.create_sheet(name)
    if proto is not None:
        _copy_header_and_layout(proto, ws)
    else:
        _write_fallback_header(ws)
    write_protocol_rows(ws, rows)
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.showGridLines = False
    return ws


def _write_fallback_header(ws: Worksheet) -> None:
    headers = [
        "Huonetila\nUtrymme",
        "Poistoilmaelin Frånluftsdon",
        "Koko Storlek",
        "Tarve Behov (l/s)",
        "Mitattu Uppmätt (l/s)",
        "Paine Tryck (Pa)",
        "Asento Justering",
        "Tuloilmaelin Tilluftsdon",
        "Koko Storlek",
        "Tarve Behov (l/s)",
        "Mitattu Uppmätt (l/s)",
        "Paine Tryck (Pa)",
        "Asento Justering",
    ]
    ws.row_dimensions[1].height = 39
    widths = [23, 14, 8.8, 7.7, 13, 13, 7.7, 14, 9.1, 7.7, 13, 13, 7.7]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, i, h)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if 2 <= i <= 7:
            cell.fill = _FILL_EX
        elif 8 <= i <= 13:
            cell.fill = _FILL_SU
        ws.column_dimensions[get_column_letter(i)].width = w


def new_workbook_from_template(template_path: Path | str | None = None) -> Workbook:
    wb = load_workbook(template_path or TEMPLATE_PATH)
    ensure_cover(wb)
    return wb


def drop_unused_prototype_sheets(wb: Workbook) -> None:
    for name in list(PROTO_SHEET_CANDIDATES):
        if name in wb.sheetnames and len(wb.sheetnames) > 1:
            del wb[name]


def save_workbook(wb: Workbook, path: Path | str | None = None) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if path is not None:
        Path(path).write_bytes(data)
    return data
