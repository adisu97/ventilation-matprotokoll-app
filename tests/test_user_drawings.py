"""Flask-flöde med syntetisk mini-ritning (inga kundfiler)."""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import load_workbook

from matprotokoll.models import ProtocolRow, SheetResult
from matprotokoll.webapp import create_app

MINI = Path(__file__).resolve().parent / "fixtures" / "mini_ritning.pdf"


def test_flask_uploads_two_synthetic_pdfs(tmp_path, monkeypatch):
    monkeypatch.setattr("matprotokoll.webapp.STORE", tmp_path)
    blob = MINI.read_bytes()

    def fake_extract(path, **kwargs):
        name = kwargs.get("pdf_name") or Path(path).name
        if "a.pdf" in name.lower():
            return SheetResult(
                pdf_name=name,
                sheet_name="plan-a",
                rows=[
                    ProtocolRow(room="E125 wc", exhaust_name="kso", exhaust_size=125),
                    ProtocolRow(room="", exhaust_name="kso", exhaust_size=125),
                ],
            )
        return SheetResult(
            pdf_name=name,
            sheet_name="plan-b",
            rows=[
                ProtocolRow(room="E102 aula", supply_name="eagle", supply_size=200),
            ],
        )

    class Dummy:
        name = "RapidOCR"

    monkeypatch.setattr("matprotokoll.webapp.extract_sheet_from_pdf", fake_extract)
    monkeypatch.setattr("matprotokoll.webapp.load_ocr_engine", lambda: Dummy())
    client = create_app().test_client()
    started = client.post(
        "/extract",
        data={
            "excel_name": "test.xlsx",
            "mode_0": "sheet",
            "mode_1": "sheet",
            "pdfs": [
                (io.BytesIO(blob), "plan-a.pdf"),
                (io.BytesIO(blob), "plan-b.pdf"),
            ],
        },
    )
    assert started.status_code == 302
    page = ""
    for _ in range(50):
        page = client.get("/status").get_data(as_text=True)
        if "Förhandsgranska" in page:
            break
    assert "plan-a" in page
    assert "plan-b" in page

    built = client.post(
        "/build",
        data={
            "excel_name": "test.xlsx",
            "nrows_0": "2",
            "s0r0c0": "E125 wc",
            "s0r0c1": "kso",
            "s0r0c2": "125",
            "s0r1c1": "kso",
            "s0r1c2": "125",
            "nrows_1": "1",
            "s1r0c0": "E102 aula",
            "s1r0c4": "eagle",
            "s1r0c5": "200",
        },
        follow_redirects=True,
    )
    assert built.status_code == 200
    assert built.data[:2] == b"PK"
    wb = load_workbook(io.BytesIO(built.data))
    assert wb.sheetnames[0] == "Koneet Maskiner"
    assert "plan-a" in wb.sheetnames
    assert "plan-b" in wb.sheetnames
    assert wb["plan-a"]["A2"].value == "E125 wc"
    assert wb["plan-a"]["A3"].value in (None, "")
    assert wb["plan-b"]["H2"].value == "eagle"
