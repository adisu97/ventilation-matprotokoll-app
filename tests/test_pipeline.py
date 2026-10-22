"""Ny-arbetsbok vs nytt blad."""

from io import BytesIO

from openpyxl import load_workbook

from matprotokoll.models import ProtocolRow, SheetResult
from matprotokoll.pipeline import build_workbooks


def _result(name: str, new_workbook: bool) -> SheetResult:
    return SheetResult(
        pdf_name=f"{name}.pdf",
        sheet_name=name,
        rows=[
            ProtocolRow(room="E125 wc", exhaust_name="kso", exhaust_size=125, supply_name="kts", supply_size=100),
            ProtocolRow(room="", exhaust_name="kso", exhaust_size=125),
        ],
        new_workbook=new_workbook,
    )


def test_same_workbook_two_sheets():
    files = build_workbooks([_result("ritning-a", False), _result("ritning-b", False)])
    assert "_base" in files
    wb = load_workbook(BytesIO(files["_base"]))
    assert "Koneet Maskiner" in wb.sheetnames
    assert "ritning-a" in wb.sheetnames
    assert "ritning-b" in wb.sheetnames
    assert wb["ritning-a"]["A2"].value == "E125 wc"
    assert wb["ritning-a"]["A3"].value in (None, "")
    assert wb["ritning-b"]["H2"].value == "kts"


def test_new_file_per_pdf():
    files = build_workbooks([_result("kok", True), _result("kallare", False)])
    assert "kok.xlsx" in files
    assert "_base" in files
    kok = load_workbook(BytesIO(files["kok.xlsx"]))
    base = load_workbook(BytesIO(files["_base"]))
    assert "kok" in kok.sheetnames
    assert "Koneet Maskiner" in kok.sheetnames
    assert "kallare" in base.sheetnames
    assert "kok" not in base.sheetnames
