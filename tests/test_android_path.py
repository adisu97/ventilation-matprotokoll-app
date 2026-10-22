import io
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from matprotokoll.models import ProtocolRow, SheetResult
from matprotokoll.ocr import TesseractEngine, _tile_is_blank, group_tesseract_lines, load_ocr_engine
from matprotokoll.pdfdoc import open_document
from matprotokoll.rooms import extract_rooms
from matprotokoll.settings import ocr_preference, ocr_rotations, quality_mode, tile_px, zoom
from matprotokoll.webapp import create_app

MINI = Path(__file__).resolve().parent / "fixtures" / "mini_ritning.pdf"


def test_settings_from_env(monkeypatch):
    monkeypatch.setenv("MATPROTOKOLL_OCR", "tesseract")
    monkeypatch.setenv("MATPROTOKOLL_ZOOM", "2")
    monkeypatch.setenv("MATPROTOKOLL_TILE", "800")
    monkeypatch.setenv("MATPROTOKOLL_QUALITY", "fast")
    monkeypatch.setenv("MATPROTOKOLL_LOWMEM", "1")
    assert ocr_preference() == "tesseract"
    assert zoom() == 2
    assert tile_px() == 800
    assert quality_mode() == "fast"
    assert ocr_rotations() == (0,)


def test_quality_accurate_defaults(monkeypatch):
    monkeypatch.delenv("MATPROTOKOLL_ZOOM", raising=False)
    monkeypatch.delenv("MATPROTOKOLL_TILE", raising=False)
    monkeypatch.delenv("MATPROTOKOLL_OVERLAP", raising=False)
    monkeypatch.delenv("MATPROTOKOLL_ROTATIONS", raising=False)
    monkeypatch.delenv("MATPROTOKOLL_LOWMEM", raising=False)
    monkeypatch.setenv("MATPROTOKOLL_QUALITY", "accurate")
    assert quality_mode() == "accurate"
    assert zoom() == 3
    assert tile_px() == 900
    assert ocr_rotations() == (0, 90, 270)


def test_ocr_preference_auto(monkeypatch):
    monkeypatch.delenv("MATPROTOKOLL_OCR", raising=False)
    assert ocr_preference() == "auto"


def test_auto_loads_rapidocr(monkeypatch):
    monkeypatch.delenv("MATPROTOKOLL_OCR", raising=False)
    engine = load_ocr_engine()
    assert engine.name == "RapidOCR"


def test_tesseract_forced_does_not_load_rapidocr(monkeypatch):
    monkeypatch.setenv("MATPROTOKOLL_OCR", "tesseract")
    engine = load_ocr_engine()
    assert engine.name == "Tesseract"


def test_group_tesseract_words_on_one_line():
    data = {
        "text": ["2xEAGLE", "CRa", "200-600", ""],
        "conf": ["90", "80", "70", "-1"],
        "left": [0, 40, 80, 0],
        "top": [0, 0, 1, 0],
        "width": [30, 20, 40, 0],
        "height": [10, 10, 10, 0],
        "block_num": [1, 1, 1, 1],
        "par_num": [1, 1, 1, 1],
        "line_num": [1, 1, 1, 1],
    }
    lines = group_tesseract_lines(data)
    assert len(lines) == 1
    assert lines[0][0] == "2xEAGLE CRa 200-600"


def test_blank_tile_without_numpy():
    assert _tile_is_blank(Image.new("RGB", (40, 40), "white"))
    assert not _tile_is_blank(Image.new("RGB", (40, 40), "black"))


@pytest.mark.skipif(shutil.which("pdftotext") is None, reason="poppler saknas")
def test_poppler_reads_room_numbers(monkeypatch):
    monkeypatch.setenv("MATPROTOKOLL_PDF", "poppler")
    with open_document(MINI) as doc:
        rooms = []
        for page in doc.pages():
            rooms.extend(extract_rooms(page))
    labels = {room.label().lower() for room in rooms}
    assert any(label.startswith("e125") for label in labels)
    assert any(label.startswith("e102") for label in labels)


@pytest.mark.skipif(shutil.which("tesseract") is None or shutil.which("pdftoppm") is None, reason="tesseract/poppler saknas")
def test_tesseract_reads_kso_from_poppler_crop(monkeypatch):
    monkeypatch.setenv("MATPROTOKOLL_PDF", "poppler")
    monkeypatch.setenv("MATPROTOKOLL_TESS_LANG", "eng")
    with open_document(MINI) as doc:
        image = doc.pages()[0].render_clip((120, 30, 230, 80), zoom=3)
    text = " ".join(hit[0] for hit in TesseractEngine().recognize(image)).upper()
    assert "KSO" in text


def test_flask_flow_names_file_then_downloads(tmp_path, monkeypatch):
    monkeypatch.setattr("matprotokoll.webapp.STORE", tmp_path)

    def fake_extract(path, **kwargs):
        return SheetResult(
            pdf_name="ritning.pdf",
            sheet_name="ritning",
            rows=[
                ProtocolRow(room="E125 wc", exhaust_name="kso", exhaust_size=125, supply_name="kts", supply_size=100),
                ProtocolRow(room="", exhaust_name="kso", exhaust_size=125),
            ],
            warnings=[],
        )

    class Dummy:
        name = "Tesseract"

    monkeypatch.setattr("matprotokoll.webapp.extract_sheet_from_pdf", fake_extract)
    monkeypatch.setattr("matprotokoll.webapp.load_ocr_engine", lambda: Dummy())
    client = create_app().test_client()
    missing = client.post("/extract", data={"excel_name": ""})
    assert missing.status_code == 400
    assert "Namn" in missing.get_data(as_text=True)

    started = client.post(
        "/extract",
        data={
            "excel_name": "telefon.xlsx",
            "mode_0": "sheet",
            "pdfs": (io.BytesIO(b"%PDF-1.4\n"), "ritning.pdf"),
        },
    )
    assert started.status_code == 302
    page = ""
    for _ in range(50):
        page = client.get("/status").get_data(as_text=True)
        if "Förhandsgranska" in page:
            break
    assert "E125 wc" in page
    assert "kso" in page

    built = client.post(
        "/build",
        data={
            "excel_name": "telefon.xlsx",
            "nrows_0": "2",
            "s0r0c0": "E125 wc",
            "s0r0c1": "kso",
            "s0r0c2": "125",
            "s0r0c4": "kts",
            "s0r0c5": "100",
            "s0r1c1": "kso",
            "s0r1c2": "125",
        },
        follow_redirects=True,
    )
    assert built.status_code == 200
    assert built.data[:2] == b"PK"
    wb = load_workbook(io.BytesIO(built.data))
    assert "Koneet Maskiner" in wb.sheetnames
    assert "ritning" in wb.sheetnames
    ws = wb["ritning"]
    assert ws["A2"].value == "E125 wc"
    assert ws["B2"].value == "kso"
    assert ws["H2"].value == "kts"
    assert ws["A3"].value in (None, "")
    assert ws["B3"].value == "kso"
