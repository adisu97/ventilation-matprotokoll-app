from io import BytesIO

from openpyxl import load_workbook

from matprotokoll.excel_io import (
    add_protocol_sheet,
    drop_unused_prototype_sheets,
    new_workbook_from_template,
    save_workbook,
)
from matprotokoll.models import ProtocolRow
from matprotokoll.sheet_names import sanitize_sheet_name, sheet_name_from_pdf, uniquify_sheet_name


def test_sanitize_invalid_chars_and_length():
    name = sheet_name_from_pdf("301 TKPK01 Vån 1-2 extra långt namn som måste kapas.pdf")
    assert len(name) <= 31
    assert ":" not in name
    bad = sanitize_sheet_name(r"a:b/c*d?e[f]g")
    assert all(ch not in bad for ch in r":/*?[]")


def test_uniquify():
    a = uniquify_sheet_name("Blad", ["Blad"])
    assert a.lower() != "blad"
    assert a.startswith("Blad")


def test_write_headers_and_continuation(tmp_path):
    rows = [
        ProtocolRow(
            room="E125 wc",
            exhaust_name="kso",
            exhaust_size=125,
            exhaust_behov=20,
            supply_name="kts",
            supply_size=125,
            supply_behov=20,
        ),
        ProtocolRow(room="", exhaust_name="kso", exhaust_size=125, exhaust_behov=20),
    ]
    wb = new_workbook_from_template()
    add_protocol_sheet(wb, "testblad", rows)
    drop_unused_prototype_sheets(wb)
    assert "Koneet Maskiner" in wb.sheetnames
    assert "testblad" in wb.sheetnames
    assert "Sheet1" not in wb.sheetnames
    ws = wb["testblad"]
    header_a = str(ws.cell(1, 1).value)
    assert "Utrymme" in header_a or "Huonetila" in header_a
    assert "Frånluftsdon" in str(ws.cell(1, 2).value)
    assert "Tilluftsdon" in str(ws.cell(1, 8).value)
    assert ws.cell(2, 1).value == "E125 wc"
    assert ws.cell(2, 2).value == "kso"
    assert ws.cell(2, 8).value == "kts"
    assert ws.cell(3, 1).value in (None, "")
    assert ws.cell(3, 2).value == "kso"
    data = save_workbook(wb, tmp_path / "out.xlsx")
    assert data[:2] == b"PK"
    loaded = load_workbook(BytesIO(data))
    assert loaded["testblad"]["A2"].value == "E125 wc"
    assert loaded["testblad"]["A3"].value in (None, "")


def test_new_file_vs_new_sheet():
    wb = new_workbook_from_template()
    add_protocol_sheet(wb, "forsta", [ProtocolRow(room="A1", exhaust_name="kso", exhaust_size=100)])
    add_protocol_sheet(wb, "andra", [ProtocolRow(room="B1", supply_name="kts", supply_size=100)])
    drop_unused_prototype_sheets(wb)
    assert "forsta" in wb.sheetnames
    assert "andra" in wb.sheetnames
    assert "Koneet Maskiner" in wb.sheetnames
